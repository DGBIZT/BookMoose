import sys
from django.core.management import BaseCommand
from django.contrib.auth import get_user_model
from genres.models import Genre
import openpyxl

User = get_user_model()

class Command(BaseCommand):
    help = 'Импортирует жанры из Excel‑файла (.xlsx)'


    def add_arguments(self, parser):
        parser.add_argument(
            'excel_file',
            type=str,
            help='Путь к Excel‑файлу (.xlsx) с жанрами'
        )
        parser.add_argument(
            '--user-id',
            type=int,
            default=1,
            help='ID пользователя, который создаст жанры (created_by)'
        )

    def handle(self, *args, **options):
        excel_file = options['excel_file']
        user_id = options['user_id']

        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            self.stderr.write(self.style.ERROR(f"Ошибка: Пользователь с ID={user_id} не найден."))
            sys.exit(1)

        try:
            workbook = openpyxl.load_workbook(excel_file, read_only=True)
            sheet = workbook.active

            if sheet.max_row < 2:
                self.stderr.write(self.style.ERROR("Ошибка: Файл пуст или нет данных."))
                sys.exit(1)

            # Читаем заголовки (первая строка)
            headers = []
            for cell in sheet[1]:
                headers.append((cell.value or '').strip())

            # Определяем индексы колонок
            try:
                name_idx = headers.index('name')
                desc_idx = headers.index('description') if 'description' in headers else -1
                order_idx = headers.index('order') if 'order' in headers else -1
                parent_idx = headers.index('parent') if 'parent' in headers else -1
            except ValueError as e:
                missing = str(e).split("'")[1]
                self.stderr.write(self.style.ERROR(f"Ошибка: Не найдено поле '{missing}' в заголовках."))
                sys.exit(1)

            created_count = 0
            updated_count = 0
            error_count = 0

            # Обрабатываем строки (начиная со 2-й)
            for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
                try:
                    # Извлекаем и приводим к строке
                    name = str(row[name_idx].value or '').strip().replace('\xa0', ' ').strip()
                    description = str(row[desc_idx].value or '').strip().replace('\xa0', ' ').strip() if desc_idx != -1 else ''
                    order_str = str(row[order_idx].value or '0').strip().replace('\xa0', ' ').strip() if order_idx != -1 else '0'
                    parent_name = str(row[parent_idx].value or '').strip().replace('\xa0', ' ').strip() if parent_idx != -1 else ''

                    # Валидация: name обязателен
                    if not name:
                        self.stderr.write(self.style.WARNING(
                            f"Строка {row_idx}: пропущена — отсутствует поле 'name'."
                        ))
                        error_count += 1
                        continue

                    # Валидация order
                    if not order_str.isdigit():
                        self.stderr.write(self.style.WARNING(
                            f"Строка {row_idx}: поле 'order' не число: {order_str}. Пропущено."
                        ))
                        error_count += 1
                        continue
                    order = int(order_str)

                    # Поиск родителя
                    parent = None
                    if parent_name:  # Только если parent_name не пустой
                        try:
                            parent = Genre.objects.get(name=parent_name)
                        except Genre.DoesNotExist:
                            self.stderr.write(self.style.WARNING(
                                f"Строка {row_idx}: родительский жанр '{parent_name}' не найден. Будет без родителя."
                            ))

                    # Создание/обновление
                    genre, created = Genre.objects.update_or_create(
                        name=name,
                        defaults={
                            'description': description,
                            'parent': parent,
                            'order': order,
                            'created_by': user,
                        }
                    )

                    if created:
                        self.stdout.write(self.style.SUCCESS(f"Создан: {name}"))
                        created_count += 1
                    else:
                        self.stdout.write(self.style.WARNING(f"Обновлён: {name}"))
                        updated_count += 1

                except Exception as e:
                    self.stderr.write(self.style.ERROR(f"Ошибка в строке {row_idx}: {e}"))
                    error_count += 1
                    continue

            # Итоговый отчёт
            self.stdout.write(self.style.SUCCESS(
                f"\nИмпорт завершён:\n"
                f"  Создано: {created_count}\n"
                f"  Обновлено: {updated_count}\n"
                f"  Ошибок: {error_count}"
            ))

        except FileNotFoundError:
            self.stderr.write(self.style.ERROR(f"Ошибка: Файл '{excel_file}' не найден."))
            sys.exit(1)
        except Exception as e:
            self.stderr.write(self.style.ERROR(f"Неожиданная ошибка: {e}"))
            sys.exit(1)
